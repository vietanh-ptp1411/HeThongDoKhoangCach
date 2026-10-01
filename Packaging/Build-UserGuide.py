"""Generate the Vietnamese operator guide as PDF and editable Word from the same content."""
from pathlib import Path
import argparse
import shutil
from xml.sax.saxutils import escape

from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser = argparse.ArgumentParser()
parser.add_argument('--screenshots', type=Path, required=True)
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
args.output.mkdir(parents=True, exist_ok=True)

P = lambda text: ('p', text)
H = lambda text: ('h', text)
N = lambda text: ('note', text)
T = lambda headers, rows, widths=None: ('table', headers, rows, widths)
I = lambda filename, caption: ('image', filename, caption)

pages = [
('HƯỚNG DẪN SỬ DỤNG', [
    ('cover', 'HỆ THỐNG ĐO\nLỰC CĂNG BELT'),
    P('PLC & Loadcell • Windows x64 • Phiên bản 1.2.1'),
    P('Tài liệu bàn giao ngày 30/09/2026 • MVA Lab'),
    N('Luồng vận hành: chọn quy cách → chọn model → scan serial → START → đo một lần → lưu kết quả và QR.'),
    H('Tra cứu nhanh'),
    T(['Nội dung', 'Trang'], [
        ['Cài đặt, mở phần mềm lần đầu', '2'], ['Thực hiện một lượt đo', '3'],
        ['Máy quét serial và các nút vận hành', '4'], ['Tìm nhanh trong Cài đặt', '5'],
        ['Danh mục quy cách và model', '6'], ['Kết nối PLC và bảng tín hiệu', '7'],
        ['QR, báo cáo và Excel', '8'], ['Dữ liệu, sao lưu và nâng cấp', '9'],
        ['Xử lý các tình huống thường gặp', '10'],
    ], [0.83, 0.17]),
    P('Đối tượng sử dụng: người vận hành máy, kỹ thuật viên cài đặt và người quản lý dữ liệu kiểm tra.'),
    P('Ảnh trong tài liệu là dữ liệu DEMO. Giá trị đo, số serial và người kiểm tra dùng để minh họa thao tác.'),
]),
('01 / Cài đặt và mở lần đầu', [
    H('Tệp bàn giao'),
    T(['Tệp', 'Công dụng'], [
        ['HeThongDoLucCangBelt-Setup-1.2.1-win-x64.exe', 'Bộ cài đặt ứng dụng và runtime .NET đi kèm.'],
        ['HuongDanSuDung.pdf', 'Hướng dẫn để đọc, in hoặc gửi cho người vận hành.'],
        ['HuongDanSuDung.docx', 'Bản Word có thể chỉnh sửa.'],
        ['SHA256SUMS.txt', 'Mã kiểm tra tính toàn vẹn của các tệp bàn giao.'],
    ], [0.55, 0.45]),
    H('Chuẩn bị'),
    P('Dùng máy Windows 11 64-bit; màn hình tối thiểu khoảng 1200 × 760 điểm hiển thị. Với Windows 10 LTSC/Enterprise, cần kiểm tra phiên bản hệ điều hành cụ thể trước khi triển khai. Bộ cài không dành cho Windows 32-bit.'),
    P('Bộ cài đã chứa .NET Desktop Runtime, không cần tải runtime riêng khi cài. Chuẩn bị kết nối mạng tới PLC và máy quét USB Keyboard/HID nếu vận hành với thiết bị thật.'),
    H('Các bước cài đặt'),
    P('1. Đóng phiên phần mềm đo đang mở, sau đó chạy tệp Setup. Bộ cài dùng giao diện Next / Install / Finish.'),
    P('2. Xác nhận tạo biểu tượng trên Desktop nếu cần. Nhấn Install, chờ giải nén và nhấn Finish.'),
    P('3. Mở “Hệ thống đo lực căng Belt” từ Desktop hoặc Start Menu. Trong Start Menu cũng có lối tắt tới bản hướng dẫn PDF.'),
    P('4. Cài đặt áp dụng cho tài khoản Windows hiện tại, tại %LOCALAPPDATA%\\Programs\\HeThongDoKhoangCach. Không cần chạy ứng dụng bằng quyền quản trị.'),
    N('Lần mở đầu tiên chạy MÔ PHỎNG (DEMO). Hãy đo thử theo trang 3 trước khi cấu hình PLC thật. Khi dùng thiết bị thật, kỹ thuật viên chuyển giao thức và kiểm tra bảng địa chỉ theo trang 7.'),
    P('Bản nội bộ này chưa được ký bằng chứng thư số của đơn vị phát hành. Nếu Windows hiện cảnh báo nhà phát hành, đối chiếu nguồn tệp và SHA256 với bên bàn giao trước khi chạy.'),
]),
('02 / Thực hiện một lượt đo', [
    I('main-window.png', 'Màn hình đo: thông tin ở trái, kết quả và QR cạnh nhau, biểu đồ và lịch sử ở phải.'),
    P('1. Kiểm tra trạng thái kết nối PLC. Có thể nhấn KẾT NỐI PLC tại thanh trạng thái; bản DEMO tự kết nối PLC mô phỏng.'),
    P('2. Chọn Quy cách. Danh sách Chủng loại chỉ hiển thị model thuộc quy cách vừa chọn.'),
    P('3. Chọn model, rồi scan serial. Khi màn hình đo đang hoạt động, có thể bấm cò máy quét ngay, không cần click ô Serial. Đợi dòng thông báo “Đã nhận serial…”.'),
    P('4. Nhấn START. Phần mềm gửi thông tin model và lệnh đo tới PLC; lựa chọn quy cách/model/serial được khóa trong khi đo.'),
    P('5. Đợi một kết quả đo hoàn tất. Màn hình hiển thị giá trị, PASS/NG và QR; mặc định kết quả được tự lưu. PLC mô phỏng hoàn tất lượt đầu sau khoảng 1,5 giây.'),
    P('6. Nếu tắt tự lưu, nhấn LƯU DỮ LIỆU. Sau đó scan lại serial để đo lượt mới; có thể scan cùng serial nếu muốn đo lại.'),
    N('Mỗi lần START chỉ nhận và lưu một kết quả. Các xung kết quả tiếp theo không tự tạo thêm dòng. Lực căng và khoảng cách tức thời vẫn cập nhật khi PLC còn kết nối.'),
]),
('03 / Máy quét và các nút vận hành', [
    H('Thiết lập máy quét serial'),
    P('Máy đọc serial sử dụng Zebra DS8178, đế thu kết nối máy tính qua USB. Cấu hình USB Keyboard HID và thêm phím Enter sau mỗi mã; phần mềm không cần chọn cổng COM.'),
    P('Trong Quick Start Guide DS8178 của Zebra (MN-002753-04EN): trang 5 hướng dẫn cắm cáp đế thu, trang 6 có mã USB Keyboard HID, trang 10 có mã Add Enter Key (Carriage Return/Line Feed). Nếu máy quét chưa ghép cặp, quét mã pairing trên đế. Tài liệu: https://www.zebra.com/content/dam/support-dam/en/documentation/unrestricted/guide/product/ds8178-qsg-en.pdf'),
    P('Thử bằng Notepad: đặt con trỏ vào trang trống, quét một mã. Serial phải xuất hiện đầy đủ và con trỏ xuống dòng. Nếu chưa được, chỉnh chế độ giao tiếp và ký tự kết thúc theo tài liệu của máy quét.'),
    P('Trong ứng dụng, chọn quy cách/model trước khi quét. Phần mềm tự nhận mã khi vừa chọn model, bấm bảng hoặc nút trên màn hình chính. Không cần nút “Kết nối máy scan” với kiểu bàn phím.'),
    N('Không nhận scan thay thế serial khi đang đo. Nếu đang nhập ô lọc ngày, mở Cài đặt/Đăng nhập/Báo cáo hoặc chuyển sang ứng dụng khác, bàn phím phục vụ cửa sổ/ô nhập đó. Trở về màn hình đo trước khi quét serial.'),
    T(['Nút', 'Chức năng'], [
        ['START', 'Đo một lần cho serial đã nhận. Sau khi hoàn tất phải scan lại cho lượt tiếp theo.'],
        ['STOP', 'Hủy lượt đang chạy. Serial được giữ để có thể nhấn START đo lại lượt bị hủy.'],
        ['RESET', 'Xóa serial, kết quả/QR trên màn hình và dữ liệu chưa lưu; giữ lựa chọn quy cách/model.'],
        ['LƯU DỮ LIỆU', 'Lưu kết quả khi tắt tự lưu hoặc thử lại sau lỗi ghi file. Không lưu trùng kết quả đã lưu.'],
        ['ĐĂNG NHẬP', 'Nhập mã người kiểm tra để gắn vào kết quả. Đây là thông tin người kiểm tra, không phải cơ chế tài khoản/mật khẩu.'],
    ], [0.27, 0.73]),
    P('USB HID là kiểu nhập liệu hiện được hỗ trợ. Phần mềm chưa có cổng nhận serial riêng qua COM/RS232 hoặc TCP/IP, và không có đèn trạng thái riêng xác nhận máy quét đã cắm.'),
]),
('04 / Tìm nhanh trong Cài đặt', [
    I('settings-section-0.png', 'Chọn nhóm ở thanh bên trái; Lưu & áp dụng và Hủy luôn nằm dưới cửa sổ.'),
    T(['Mục', 'Nội dung'], [
        ['Kết nối PLC', 'Chọn mô phỏng / MC 3E / Modbus, IP, port, Unit ID, chu kỳ đọc, timeout và thử kết nối.'],
        ['Địa chỉ PLC', 'Sửa tag lệnh, lực căng, khoảng cách, kết quả, trạng thái; chọn kiểu dữ liệu, hệ số và thứ tự word.'],
        ['Quy cách & model', 'Hai bảng cạnh nhau để sửa giới hạn đo, điều kiện, model và bit PLC.'],
        ['QR kết quả', 'Đặt mẫu nội dung mã QR và xem trước nội dung mẫu.'],
        ['Thông tin & dữ liệu', 'Tiêu đề, đơn vị, người kiểm tra, file kết quả, số cột biểu đồ và tự lưu.'],
    ], [0.28, 0.72]),
    P('Nhấn Lưu & áp dụng để lưu các mục đã chỉnh. Hủy bỏ thay đổi chưa lưu. Nếu có lỗi nhập liệu, màn hình chuyển tới nhóm cần sửa. Cài đặt chỉ mở được khi lượt đo đã kết thúc và không còn kết quả chờ lưu.'),
]),
('05 / Quy cách và model', [
    I('settings-section-2.png', 'Chọn quy cách ở bảng trái; bảng phải hiển thị các model thuộc quy cách đó.'),
    T(['Quy cách', 'Model', 'Bit MC mặc định'], [
        ['3 ~ 4 mm', 'LCP', 'M45'],
        ['4 ~ 5 mm', 'M1 / M2 / NF / PP1', 'M42 / M42 / M43 / M41'],
        ['3.5 ~ 6 mm', 'GSM / CPX', 'M40 / M44'],
    ], [0.23, 0.36, 0.41]),
    P('1. Chọn một dòng quy cách, sửa LSL (giới hạn dưới), USL (giới hạn trên) và điều kiện. LSL phải nhỏ hơn USL; đơn vị là mm.'),
    P('2. Sửa hoặc thêm model ở bảng bên phải. Bit PLC là địa chỉ phần mềm bật khi chọn model; các bit model khác được tắt.'),
    P('3. Nhấn Lưu & áp dụng. M1 và M2 được dùng chung M42 trong cùng quy cách; không dùng một bit cho hai quy cách khác nhau.'),
    N('Bảng trên là cấu hình mặc định theo tài liệu đầu vào, điều kiện “Tác động lực 2gf”. Kỹ thuật viên phải đối chiếu bit với chương trình PLC thực tế, đặc biệt LCP / M45. Modbus dùng coil C tương ứng.'),
    P('Danh mục lưu trong Data\\HeThongDo.db. Thay đổi quy cách/model không sửa lại giới hạn hoặc QR đã lưu trong các kết quả lịch sử.'),
]),
('06 / Kết nối PLC và tín hiệu', [
    P('Vào Cài đặt → Kết nối PLC, chọn giao thức, nhập IP/port và nhấn Kiểm tra kết nối. Sau đó vào Địa chỉ PLC để đối chiếu từng tag. MC dùng 3E frame, Binary qua TCP; Modbus dùng TCP và Unit ID.'),
    P('Nút Mặc định MC / Mặc định Modbus khôi phục bảng địa chỉ mặc định. Chỉ dùng khi phù hợp với chương trình PLC tại máy; cấu hình mặc định chưa chứng minh đã khớp thiết bị thực tế.'),
    T(['Tín hiệu', 'MC', 'Modbus', 'Kiểu mặc định'], [
        ['START / STOP / RESET', 'M0 / M15 / M16', 'C0 / C15 / C16', 'Bit'],
        ['Lực căng', 'D100', 'HR100', 'Float32'],
        ['Khoảng cách', 'D102', 'HR102', 'Float32'],
        ['Kết quả đo chốt', 'D104', 'HR104', 'Float32'],
        ['TOTAL / OK / NG', 'D110 / D112 / D114', 'HR110 / HR112 / HR114', 'Int32'],
        ['Loadcell ổn định', 'M100', 'C100', 'Bit'],
        ['Đo xong', 'M101', 'C101', 'Bit'],
        ['Phân định OK / NG', 'M102 / M103', 'C102 / C103', 'Bit'],
        ['Máy đang chạy', 'M104', 'C104', 'Bit'],
        ['LSL / USL ghi xuống', 'Để trống', 'Để trống', 'Float32, tùy chọn'],
    ], [0.29, 0.23, 0.30, 0.18]),
    H('Trình tự phía PLC'),
    P('Mỗi START thực hiện một lần đo. PLC ghi xong giá trị, OK/NG và bộ đếm rồi mới bật Đo xong; giữ bit ít nhất một chu kỳ đọc (nên ≥ 500 ms), hạ bit trước lượt kế tiếp. Hạ bit Máy đang chạy khi hoàn tất. PLC tự xóa bit lệnh sau khi xử lý.'),
    P('Để trống Đo xong: phần mềm theo TOTAL tăng. Nếu cả hai trống: theo giá trị kết quả thay đổi; cách này không phân biệt được hai lượt có cùng giá trị, vì vậy ưu tiên bit Đo xong hoặc TOTAL.'),
    P('Để trống cả OK và NG: phần mềm tự so giá trị với LSL/USL (bao gồm hai biên). Để trống kết quả chốt: lấy khoảng cách tại lúc Đo xong. Để trống bộ đếm: đếm theo lịch sử đã lọc.'),
    N('Kiểu dữ liệu, hệ số và thứ tự word phải khớp PLC. Ví dụ PLC lưu 370 để biểu diễn 3.70 mm thì hệ số là 0.01. Kết quả 0 vẫn có thể là một kết quả hợp lệ khi PLC đã báo đo xong.'),
]),
('07 / QR, báo cáo và Excel', [
    I('report-qr.png', 'BÁO CÁO có cột QR lưu cùng từng kết quả; rê chuột trên mã để xem phóng to.'),
    P('Mặc định QR chứa giá trị như 3.70. Trong Cài đặt → QR kết quả, có thể dùng {Serial};{Model};{Value};{Result}. Ví dụ: 985X57200-00123;CPX;3.70;OK. Ảnh DEMO trong tài liệu dùng mẫu ghép để minh họa.'),
    P('QR chỉ xuất hiện khi có kết quả hoàn chỉnh. Mã được lưu tại thời điểm đo; đổi mẫu sau này không thay đổi QR cũ. Scan lượt mới hoặc RESET xóa QR trên màn hình đo nhưng không xóa bản ghi đã lưu.'),
    P('Mở BÁO CÁO, chọn khoảng ngày, model, OK/NG và từ khóa rồi nhấn Tìm. Từ khóa tìm được trong serial, nội dung QR, người kiểm tra và ghi chú. Xóa lọc để trở lại toàn bộ dữ liệu.'),
    P('Nhấn Excel để xuất những dòng đã lọc. File .xlsx có các thông tin đo và cột QR gồm ảnh mã cùng nội dung chữ. Dữ liệu cũ chưa có QR thì để trống cột này; phần mềm không suy đoán lại mã cũ.'),
    P('Nút XUẤT EXCEL ở màn hình chính xuất theo bộ lọc ngày hiện tại. THỐNG KÊ tổng hợp theo model/giới hạn lịch sử; histogram trên màn hình chính gộp kết quả theo quy cách đã lưu.'),
    N('Máy quét nhận serial và máy quét đọc QR kết quả là hai thao tác khác nhau. Cần thử quét mã trên màn hình thực tế và xác nhận nội dung mà hệ thống phía Brother yêu cầu trước khi bàn giao vận hành.'),
]),
('08 / Dữ liệu, sao lưu và nâng cấp', [
    P('Thư mục mặc định của bản cài: %LOCALAPPDATA%\\Programs\\HeThongDoKhoangCach. Có thể dán đường dẫn này vào thanh địa chỉ File Explorer để mở.'),
    T(['Tệp / thư mục', 'Nội dung'], [
        ['appsettings.json', 'Kết nối, địa chỉ PLC, mẫu QR, đường dẫn và các lựa chọn gần nhất.'],
        ['Data\\HeThongDo.db', 'Database SQLite chứa danh mục quy cách/model.'],
        ['Data\\results.jsonl', 'Lịch sử đo; mỗi dòng một kết quả, gồm QrText đã lưu.'],
        ['HuongDan', 'Bản PDF và Word của tài liệu này.'],
    ], [0.37, 0.63]),
    H('Sao lưu / phục hồi'),
    P('Đóng ứng dụng, sao chép appsettings.json và toàn bộ thư mục Data sang vị trí sao lưu có ghi ngày. Nếu đã cấu hình đường dẫn khác hoặc thư mục mạng, sao lưu cả các tệp ở vị trí đó.'),
    P('Để phục hồi: đóng phần mềm, sao lưu bộ dữ liệu hiện tại, chép bộ đã lưu về đúng vị trí rồi mở lại. Việc phục hồi có thể thay thế dữ liệu mới hơn, cần chọn đúng thời điểm sao lưu.'),
    H('Nâng cấp / gỡ cài đặt'),
    P('Bản 1.2.1 không dùng file master, Đơn hàng hoặc Line. Lịch sử và QR đã lưu từ bản trước vẫn đọc được. Nếu mẫu QR cũ có trường {OrderNo} hoặc {Line}, mẫu dùng cho lượt đo mới được đưa về {Value}; có thể chỉnh lại trong Cài đặt → QR kết quả.'),
    P('Đóng ứng dụng, chạy bộ cài mới bằng cùng tài khoản Windows. Bộ cài giữ appsettings.json, danh mục và lịch sử hiện có. Gỡ bằng Windows Settings → Apps hoặc lối tắt Gỡ cài đặt; dữ liệu người dùng được giữ để sao lưu hay cài lại.'),
    N('Không sửa tay results.jsonl khi phần mềm đang chạy. Khi chuyển máy hoặc tài khoản Windows, cần chép dữ liệu sang thư mục tương ứng; dữ liệu mặc định thuộc từng tài khoản Windows.'),
]),
('09 / Xử lý tình huống thường gặp', [
    T(['Hiện tượng', 'Cách xử lý'], [
        ['START chưa bấm được', 'Kiểm tra đã chọn quy cách/model, scan xong serial, PLC ONLINE và lượt trước đã lưu. Sau mỗi kết quả phải scan lại cho lượt mới.'],
        ['Scan không nhận serial', 'Thử Notepad; kiểm tra USB Keyboard/HID và Enter. Trở về màn hình đo, đóng hộp thoại, chọn model. Nếu đang sửa ô ngày, kết thúc thao tác lọc trước khi scan.'],
        ['Nhận mã thiếu/sai ký tự', 'Kiểm tra chế độ bàn phím của máy quét, bố cục bàn phím và bộ gõ trên máy tính. Đối chiếu nội dung quét trong Notepad với nhãn thực tế.'],
        ['PLC OFFLINE', 'Kiểm tra nguồn PLC, cáp mạng, IP/port/Unit ID và giao thức. Dùng Kiểm tra kết nối; nhờ kỹ thuật viên đối chiếu cấu hình Ethernet của PLC.'],
        ['Đang đo nhưng không ra kết quả', 'Kiểm tra bit Đo xong hoặc TOTAL, giá trị chốt và trạng thái chạy. PLC phải phát tín hiệu hoàn tất đúng trình tự ở trang 7.'],
        ['Chờ OK/NG', 'PLC đã báo đo xong nhưng chưa có bit phân định. Kiểm tra địa chỉ OK/NG; nếu muốn phần mềm tự so quy cách, kỹ thuật viên để trống cả hai tag.'],
        ['Không lưu được kết quả', 'Kiểm tra đường dẫn, quyền ghi và dung lượng ổ đĩa. Kết quả chưa ghi còn giữ trong bộ nhớ; sửa nguyên nhân rồi bấm LƯU DỮ LIỆU, không RESET hoặc đóng ứng dụng trước khi lưu.'],
        ['QR không có / khó quét', 'Đợi kết quả hoàn chỉnh, kiểm tra mẫu QR. Bản ghi cũ chưa lưu QR sẽ trống. Trong Báo cáo rê chuột để phóng to; kiểm tra độ sáng và khoảng cách quét.'],
        ['Không thấy kết quả trong báo cáo', 'Kiểm tra bộ lọc ngày, model, OK/NG và từ khóa; bấm Xóa lọc. Kiểm tra đúng đường dẫn file lịch sử.'],
    ], [0.29, 0.71]),
    H('Thông tin cần gửi khi báo lỗi'),
    P('Gửi ảnh màn hình/thông báo lỗi, phiên bản 1.2.1, thời điểm xảy ra, serial, giao thức PLC và các bước vừa thao tác. Với lỗi dữ liệu, giữ bản sao cấu hình và lịch sử liên quan để kỹ thuật viên đối chiếu.'),
    P('Tài liệu tham khảo nền tảng: https://learn.microsoft.com/dotnet/core/install/windows • Công cụ đóng gói: NSIS (https://nsis.sourceforge.io). Hướng dẫn nghiệp vụ và ảnh được đối chiếu với mã nguồn bản 1.2.1.'),
]),
]

fontdir = Path('C:/Windows/Fonts')
for name, file in [('Guide', 'arial.ttf'), ('GuideBold', 'arialbd.ttf'), ('GuideItalic', 'ariali.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(fontdir/file)))
pdfmetrics.registerFontFamily('Guide', normal='Guide', bold='GuideBold', italic='GuideItalic', boldItalic='GuideBold')
navy = colors.HexColor('#183D65')
blue = colors.HexColor('#235FA4')
light = colors.HexColor('#EAF1F9')
styles = getSampleStyleSheet()
styles.add(ParagraphStyle('BodyVI', fontName='Guide', fontSize=10.1, leading=14.1, spaceAfter=7, textColor=colors.HexColor('#263445')))
styles.add(ParagraphStyle('TitleVI', fontName='GuideBold', fontSize=21, leading=27, textColor=navy, spaceAfter=16))
styles.add(ParagraphStyle('CoverVI', fontName='GuideBold', fontSize=30, leading=38, textColor=navy, spaceBefore=18, spaceAfter=20))
styles.add(ParagraphStyle('HeadingVI', fontName='GuideBold', fontSize=12.5, leading=17, textColor=blue, spaceBefore=9, spaceAfter=7))
styles.add(ParagraphStyle('CellVI', parent=styles['BodyVI'], fontSize=9, leading=12, spaceAfter=0))
styles.add(ParagraphStyle('HeadCellVI', parent=styles['CellVI'], fontName='GuideBold', textColor=colors.white))
styles.add(ParagraphStyle('CaptionVI', parent=styles['BodyVI'], fontSize=8.5, leading=11, textColor=colors.HexColor('#5A6B7F'), spaceAfter=12))
styles.add(ParagraphStyle('NoteVI', parent=styles['BodyVI'], fontSize=10, leading=14, spaceAfter=0))

def para(text, style='BodyVI'):
    return Paragraph(escape(text).replace('\n', '<br/>'), styles[style])

page_width, page_height = A4
usable = page_width - 84

def footer(canvas, doc):
    canvas.setTitle('Hướng dẫn sử dụng - Hệ thống đo lực căng Belt 1.2.1')
    canvas.setAuthor('MVA Lab')
    canvas.setStrokeColor(colors.HexColor('#D4DFEA'))
    canvas.line(42, 36, page_width-42, 36)
    canvas.setFont('Guide', 8)
    canvas.setFillColor(colors.HexColor('#64748B'))
    canvas.drawString(42, 24, 'MVA Lab  •  Hệ thống đo lực căng Belt  •  v1.2.1')
    canvas.drawRightString(page_width-42, 24, f'Trang {doc.page}')
    if doc.page > 1:
        canvas.setFont('Guide', 8)
        canvas.drawString(42, page_height-24, 'HƯỚNG DẪN SỬ DỤNG / 30.09.2026')

story=[]
word=Document()
section=word.sections[0]
section.page_width=Cm(21); section.page_height=Cm(29.7)
section.top_margin=Cm(1.55); section.bottom_margin=Cm(1.65)
section.left_margin=Cm(1.55); section.right_margin=Cm(1.55)
normal=word.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.1
for style in ['Title','Heading 1','Heading 2']:
    word.styles[style].font.name='Arial'; word.styles[style].font.color.rgb=RGBColor.from_string('183D65')
word.styles['Heading 1'].font.size=Pt(20)
word.styles['Heading 2'].font.size=Pt(12)
word.core_properties.title='Hướng dẫn sử dụng - Hệ thống đo lực căng Belt 1.2.1'
word.core_properties.author='MVA Lab'
wp=section.footer.paragraphs[0]
wp.add_run('MVA Lab • v1.2.1     |     Trang ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); wp._p.append(field)
for run in wp.runs: run.font.size=Pt(8)

for page_index,(title,blocks) in enumerate(pages):
    if page_index: story.append(PageBreak()); word.add_page_break()
    story.append(para(title,'TitleVI')); word.add_heading(title,level=1)
    for block in blocks:
        kind=block[0]
        if kind in ('p','h','cover'):
            style={'p':'BodyVI','h':'HeadingVI','cover':'CoverVI'}[kind]
            story.append(para(block[1],style))
            if kind=='h': word.add_heading(block[1],level=2)
            elif kind=='cover':
                pcover=word.add_paragraph();r=pcover.add_run(block[1]);r.bold=True;r.font.size=Pt(30);r.font.color.rgb=RGBColor.from_string('183D65')
            else: word.add_paragraph(block[1])
        elif kind=='note':
            box=Table([[para(block[1],'NoteVI')]],colWidths=[usable])
            box.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),light),('BOX',(0,0),(-1,-1),0.5,colors.HexColor('#C8D9EE')),('LEFTPADDING',(0,0),(-1,-1),12),('RIGHTPADDING',(0,0),(-1,-1),12),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]))
            story.extend([box,Spacer(1,10)])
            pn=word.add_paragraph();pn.add_run(block[1]).bold=True
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF1F9');pn._p.get_or_add_pPr().append(shade)
        elif kind=='table':
            headers,rows,widths=block[1:]
            widths=widths or [1/len(headers)]*len(headers)
            data=[[para(x,'HeadCellVI') for x in headers]]+[[para(x,'CellVI') for x in row] for row in rows]
            table=Table(data,colWidths=[usable*w for w in widths],repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),navy),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FC')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),0.35,colors.HexColor('#D6DFEA')),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story.extend([table,Spacer(1,10)])
            wt=word.add_table(rows=1,cols=len(headers));wt.style='Light Shading Accent 1'
            for c,text in zip(wt.rows[0].cells,headers): c.text=text
            for row in rows:
                for c,text in zip(wt.add_row().cells,row): c.text=text
            for row in wt.rows:
                for c,w in zip(row.cells,widths):
                    c.width=Cm(17.9*w)
                    for paragraph in c.paragraphs:
                        for run in paragraph.runs:run.font.size=Pt(9)
            word.add_paragraph()
        elif kind=='image':
            source=args.screenshots/block[1]
            if not source.exists(): raise FileNotFoundError(source)
            picture=Image(str(source))
            scale=min(usable/picture.imageWidth,245/picture.imageHeight)
            picture.drawWidth=picture.imageWidth*scale;picture.drawHeight=picture.imageHeight*scale
            story.append(KeepTogether([picture,Spacer(1,6),para(block[2],'CaptionVI')]))
            word.add_picture(str(source),width=Cm(min(17.9,picture.drawWidth/72*2.54)))
            wp=word.add_paragraph(block[2]);wp.style='Caption'

pdf=SimpleDocTemplate(str(args.output/'HuongDanSuDung.pdf'),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=43,bottomMargin=48)
pdf.build(story,onFirstPage=footer,onLaterPages=footer)
word.save(args.output/'HuongDanSuDung.docx')
print(args.output/'HuongDanSuDung.pdf')
print(args.output/'HuongDanSuDung.docx')
